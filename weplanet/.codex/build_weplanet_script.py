from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = r"C:\hwiwhi\Github\project-DSA\weplanet\WePlaNet_발표대본_4분.docx"


slides = [
    ("슬라이드 1  표지", "0:00–0:15", "안녕하세요. 팬과 아티스트를 잇는 통합 팬 커뮤니티 웹 서비스, WePlaNet을 개발한 IT 4조입니다. 지금부터 프로젝트의 기획 배경과 핵심 기능, 그리고 개발 과정에서 얻은 경험을 소개하겠습니다."),
    ("슬라이드 2  기획 배경", "0:15–0:35", "기존 팬 활동은 게시판, 메신저, 일정 안내, 공동 구매와 굿즈 결제처럼 여러 서비스에 나뉘어 있습니다. 사용자는 계속 서비스를 옮겨 다녀야 하고, 운영자는 공식 글 구분이나 악성 채팅 대응에 어려움이 있습니다. WePlaNet은 이런 기능을 아티스트별 커뮤니티 안에 모아 팬에게는 이어지는 참여 경험을, 운영자에게는 체계적인 관리 도구를 제공하고자 했습니다."),
    ("슬라이드 3  서비스 구조", "0:35–0:58", "서비스는 팬, 아티스트, 소속사, 관리자 역할로 나뉩니다. 팬은 커뮤니티 활동과 구매 기능을 이용하고, 아티스트는 그룹 로그인 뒤 멤버 프로필을 선택해 공식 글과 DM을 관리합니다. 소속사는 포털에서 일정, 미디어, 굿즈와 멤버를 운영하고, 관리자는 입점 승인과 신고 제재, 프로젝트 심사와 정산을 담당합니다. 역할마다 화면과 권한을 분리한 것이 구조의 핵심입니다."),
    ("슬라이드 4  주요 기능", "0:58–1:15", "주요 기능은 커뮤니티 활동, 실시간 소통, 참여와 결제, 아티스트 포털, 관리자 기능으로 구성했습니다. 게시글을 보는 데서 끝나지 않고 DM과 라이브, 팬 프로젝트와 멤버십, 굿즈 구매까지 하나의 서비스 안에서 이어지도록 설계했습니다. 여기에 소셜 로그인과 다국어, AI 번역과 요약 기능도 적용했습니다."),
    ("슬라이드 5  기술 구성", "1:15–1:32", "화면은 Thymeleaf 기반 서버 렌더링과 자바스크립트 비동기 처리를 함께 사용했습니다. 서버는 Java 21과 Spring Boot, Security, JPA로 구성했고 데이터는 MySQL에 저장했습니다. 실시간 기능에는 WebSocket, STOMP, WebRTC를 사용했으며, 결제와 AI 기능은 Toss Payments와 Gemini를 연동했습니다."),
    ("슬라이드 6  팬 프로젝트와 결제", "1:32–2:08", "첫 번째 집중 개발 기능은 팬 프로젝트와 결제입니다. 실제 돈이 모이는 기능인 만큼 아무나 프로젝트를 열 수 없도록 했습니다. 일반 배지 5개와 스페셜 배지 1개, 가입 이메일 본인 확인, 정산 계좌 암호화라는 세 단계 조건을 두고 관리자 승인 이후에만 모금이 시작됩니다. 결제는 준비, 입금 대기, 결제 완료 상태로 관리합니다. 더블 클릭으로 같은 주문이 생기지 않도록 멱등 키를 사용했고, 결제 금액도 서버의 데이터와 다시 비교했습니다. 입금 확인은 웹훅과 5분 주기 확인을 함께 사용하고, 주문 행을 잠가 배지가 중복 지급되지 않도록 했습니다."),
    ("슬라이드 7  해시태그 총공 이벤트", "2:08–2:43", "두 번째는 해시태그 총공 이벤트입니다. 외부 SNS API는 비용이 들고 참여 팬덤을 정확히 구분하기 어려워, 사이트 내부 팬 게시판 글을 기준으로 집계했습니다. 관리자가 기간과 참여 팀, 해시태그를 등록하면 글 저장 뒤 자동으로 유효 여부를 판정합니다. 아티스트에게 숨긴 글, 커뮤니티 미가입자, 하루 세 건을 초과한 글은 제외 사유와 함께 기록됩니다. 순위는 단순 글 수가 아니라 전체 가입자 중 참여한 사람의 비율을 사용해 작은 팬덤도 경쟁할 수 있게 했고, 종료 후에는 결과를 고정해 공지 초안까지 자동 생성합니다."),
    ("슬라이드 8  실시간 소통과 안전한 운영", "2:43–3:10", "세 번째는 실시간 소통과 운영 안전 기능입니다. 1대1 DM은 STOMP로 실시간 송수신하고, 팬이 아티스트에게 보내는 메시지는 하루 10회로 제한했습니다. 라이브 방송은 WebRTC로 영상을 송출하고 STOMP로 채팅을 처리합니다. 소통 기능과 함께 금칙어 필터, 신고와 제재, 관리자 이메일 2단계 인증을 적용했으며, 관리자의 모든 조치는 로그로 남도록 했습니다."),
    ("슬라이드 9  어려웠던 점", "3:10–3:39", "가장 어려웠던 점은 여섯 명의 작업을 하나의 서비스로 합치는 과정이었습니다. 87개의 기능 브랜치를 병합하면서 기존 기능이 사라지는 회귀가 발생해 백업 브랜치와 커밋 규칙을 정리했습니다. 또 팀원마다 데이터베이스 상태가 달라 서버가 실행되지 않는 문제가 있어 초기화 SQL과 증분 SQL을 분리하고 테이블을 간소화했습니다. 역할이 늘면서 권한 검사도 복잡해졌기 때문에, 본인 커뮤니티를 판정하는 로직을 공통화해 반복 오류를 줄였습니다."),
    ("슬라이드 10  향후 개발", "3:39–4:00", "앞으로는 프로젝트가 무산됐을 때의 결제 취소와 환불, 정산 계좌 실명 확인을 보완할 계획입니다. 모바일 반응형도 포털과 관리자 화면까지 확대하고, 총공 우승 팬덤 보상과 외부 SNS 집계를 추가하고자 합니다. 아티스트 멤버별 팔로우와 프로필 사진, 로그인 시도 제한도 개선 과제로 남아 있습니다."),
    ("슬라이드 11  마무리", "4:00–4:12", "WePlaNet을 개발하며 중요한 데이터는 서버에서 다시 검증하고, 동시 처리는 잠금으로 막으며, 운영자의 조치는 기록으로 남겨야 한다는 점을 배웠습니다. 이상으로 발표를 마치겠습니다. 감사합니다."),
]


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)

styles = doc.styles
styles["Normal"].font.name = "맑은 고딕"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(32, 32, 32)

title_style = styles["Title"]
title_style.font.name = "맑은 고딕"
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

heading_style = styles["Heading 1"]
heading_style.font.name = "맑은 고딕"
heading_style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
heading_style.font.size = Pt(13)
heading_style.font.bold = True
heading_style.font.color.rgb = RGBColor(0, 0, 0)
heading_style.paragraph_format.space_before = Pt(10)
heading_style.paragraph_format.space_after = Pt(3)
heading_style.paragraph_format.keep_with_next = True

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("WePlaNet 프로젝트 발표 대본")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("발표 목표 시간 약 4분 10초  |  최대 5분")
r.bold = True
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(64, 74, 97)
p.paragraph_format.space_after = Pt(8)

intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
intro.add_run("슬라이드가 바뀐 뒤 한 박자 쉬고 읽습니다. 굵게 표시된 구간명은 읽지 않습니다.")
intro.runs[0].font.size = Pt(9.5)
intro.runs[0].font.color.rgb = RGBColor(90, 90, 90)
intro.paragraph_format.space_after = Pt(12)

for idx, (heading, timing, script) in enumerate(slides, 1):
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(heading)
    t = p.add_run(f"   {timing}")
    t.font.size = Pt(10)
    t.font.bold = False
    t.font.color.rgb = RGBColor(64, 91, 166)

    p = doc.add_paragraph(script)
    p.paragraph_format.line_spacing = 1.22
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True

    if idx == 5:
        doc.add_page_break()

doc.add_page_break()
p = doc.add_paragraph(style="Heading 1")
p.add_run("발표 직전 체크")
checks = [
    "슬라이드 6부터 말이 빨라지지 않도록 결제 상태 세 단어를 또렷하게 말하기",
    "슬라이드 7의 참여율 계산은 ‘참여 가입자 나누기 전체 가입자’라고 그대로 설명하기",
    "시간이 부족하면 슬라이드 4의 마지막 문장과 슬라이드 10의 세부 항목을 생략하기",
    "마지막 문장을 말한 뒤 화면을 보고 ‘감사합니다’로 끝내기",
]
for item in checks:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item)
    p.paragraph_format.space_after = Pt(4)

doc.core_properties.title = "WePlaNet 프로젝트 발표 대본"
doc.core_properties.subject = "약 4분 분량 발표 대본"
doc.core_properties.author = "IT 4조"
doc.save(OUTPUT)
print(OUTPUT)
